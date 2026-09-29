"""Gerçek model sonuçlarından Türkçe proje raporunu oluştur."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


KOK = Path(__file__).resolve().parent
CIKTI_KLASORU = KOK / "outputs"
istatistikler = json.loads((CIKTI_KLASORU / "results.json").read_text(encoding="utf-8"))
modeller = istatistikler["models"]
temel_model = modeller["Çoğunluk modeli"]["test"]
lojistik_model = modeller["Lojistik regresyon"]["test"]
orman_modeli = modeller["Random Forest"]["test"]
duyarlilik_sonucu = istatistikler["random_forest_without_page_values"]["test"]
hata_profilleri = istatistikler["error_profiles"]


KOD_DEPOSU = "https://github.com/kucukenes17/e-ticaret-satin-alma-tahmini"


def yuzde(deger: float) -> str:
    """Ondalık başarı değerini Türkçe yüzde yazımına çevir."""
    return f"%{deger * 100:.1f}".replace(".", ",")


def zemin_rengi(cell, fill: str):
    """Tablo hücresinin arka plan rengini ayarla."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def kenarliklar(cell):
    """Tablo hücresine açık gri kenarlık ekle."""
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        node.append(item)
    tc_pr.append(node)


rapor = Document()
# Dört sayfalık raporun ortak sayfa ve yazı düzeni.
sayfa_bolumu = rapor.sections[0]
sayfa_bolumu.page_width = Inches(8.5)
sayfa_bolumu.page_height = Inches(11)
sayfa_bolumu.top_margin = Inches(0.68)
sayfa_bolumu.bottom_margin = Inches(0.63)
sayfa_bolumu.left_margin = Inches(0.76)
sayfa_bolumu.right_margin = Inches(0.76)
sayfa_bolumu.header_distance = Inches(0.34)
sayfa_bolumu.footer_distance = Inches(0.36)

stiller = rapor.styles
normal_stil = stiller["Normal"]
normal_stil.font.name = "Arial"
normal_stil.font.size = Pt(10.3)
normal_stil.font.color.rgb = RGBColor(0, 0, 0)
normal_stil.paragraph_format.space_after = Pt(5)
normal_stil.paragraph_format.line_spacing = 1.13
for name, size, before, after in [
    ("Title", 18, 0, 7),
    ("Heading 1", 12.8, 11, 4),
    ("Heading 2", 10.8, 8, 3),
]:
    style = stiller[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
title_ppr = stiller["Title"]._element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)

ustbilgi = sayfa_bolumu.header.paragraphs[0]
ustbilgi.text = ""
ustbilgi.style = normal_stil
ustbilgi.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in ustbilgi.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(80, 80, 80)

altbilgi = sayfa_bolumu.footer.paragraphs[0]
altbilgi.alignment = WD_ALIGN_PARAGRAPH.CENTER
altbilgi.add_run("Sayfa ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
altbilgi._p.append(field)
for run in altbilgi.runs:
    run.font.size = Pt(8)


def paragraf(text: str = "", bold_lead: str | None = None):
    """Raporun gövdesine normal bir paragraf ekle."""
    p = rapor.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    return p


def baslik(text: str, level: int = 1):
    """Belirtilen düzeyde bölüm başlığı ekle."""
    return rapor.add_heading(text, level)


def gorsel(filename: str, caption: str, width=5.65):
    """Grafiği ve altındaki Türkçe açıklamayı ekle."""
    p = rapor.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(CIKTI_KLASORU / filename), width=Inches(width))
    cap = rapor.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    for run in cap.runs:
        run.font.size = Pt(8.4)
        run.font.italic = True


def metrik_tablosu():
    """Üç modelin test sonuçlarını karşılaştıran tabloyu oluştur."""
    values = [
        ("Çoğunluk sınıfı", temel_model),
        ("Lojistik regresyon", lojistik_model),
        ("Random Forest", orman_modeli),
    ]
    tablo = rapor.add_table(rows=1, cols=5)
    tablo.autofit = False
    widths = [2.0, 1.1, 1.1, 1.1, 1.1]
    for cell, label, width in zip(tablo.rows[0].cells, ["Model", "Accuracy", "Precision", "Recall", "F1"], widths):
        cell.text = label
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        zemin_rengi(cell, "DCE8F2")
        kenarliklar(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    for i, (name, item) in enumerate(values):
        cells = tablo.add_row().cells
        vals = [name, yuzde(item["accuracy"]), yuzde(item["precision"]), yuzde(item["recall"]), yuzde(item["f1"])]
        for cell, value, width in zip(cells, vals, widths):
            cell.text = value
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            kenarliklar(cell)
            if i % 2:
                zemin_rengi(cell, "F6F9FC")
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.space_before = Pt(2)
                if cell != cells[0]:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(9)
    return tablo


rapor.add_paragraph("E-ticaret oturumlarında satın alma tahmini", style="Title")
paragraf("UCI Online Shoppers Purchasing Intention veri seti üzerinde sınıflandırma çalışması")
paragraf(
    "Özet. 12.330 oturumluk açık veri setinde satın alma ile sonuçlanan oturumları tahmin etmek için "
    "lojistik regresyon ve Random Forest karşılaştırıldı. Tamamen aynı 125 kayıt çıkarıldıktan sonra "
    f"{istatistikler['analysis_rows']:,} satır analiz edildi. Doğrulama F1 puanına göre seçilen Random Forest, "
    f"ayrı test kümesinde {yuzde(orman_modeli['recall'])} recall ve {yuzde(orman_modeli['f1'])} F1 elde etti. "
    "Sonuçlar özellikle PageValues değişkeninin erişilebilirliğine duyarlıdır."
)

baslik("1 Problem ve veri seti")
paragraf(
    "Bir e-ticaret oturumunun satın alma ile bitip bitmediğini oturum özelliklerinden sınıflandırmak "
    "amaçlanmıştır. Girdiler ziyaret edilen sayfaların türü ve süresi, trafik ve ziyaretçi özellikleridir; "
    "hedef çıktı Revenue değişkenidir (satın alma: evet/hayır). Böyle bir model, oturumların "
    "davranış örüntülerini anlamaya yardımcı olabilir. Gerçek zamanlı kullanım için her özelliğin "
    "tahmin anında hazır olduğunun ayrıca doğrulanması gerekir."
)
paragraf(
    "Veri kaynağı UCI Machine Learning Repository'dir (Sakar ve Kastro, 2018). Ham veri 12.330 satır, "
    "17 girdi değişkeni ve 1 hedef değişken içerir. Eksik hücre yoktur. Tamamen aynı 125 satır, "
    "eğitim ve test kümeleri arasında birebir kopya bulunmasını önlemek amacıyla çıkarılmıştır; "
    "kimlik alanı olmadığı için bunların ayrı gerçek oturumlar olma ihtimali bir sınırlılıktır."
)
gorsel("class_balance.png", "Şekil 1. Satın alma sınıfı, tekrar eden kayıtlar çıkarıldıktan sonra.", 5.65)
paragraf(
    f"Temizlenen veride satın alma oranı {yuzde(istatistikler['positive_rate'])}. Bu dengesizlik nedeniyle "
    "çoğunluk sınıfını sürekli tahmin eden bir model yüksek accuracy gösterse bile satın almaları kaçırabilir."
)

rapor.add_page_break()
baslik("2 Veri keşfi ve ön işleme")
gorsel("visitor_type.png", "Şekil 2. Ziyaretçi türüne göre gözlenen satın alma oranı.", 5.4)
paragraf(
    f"Yeni ziyaretçilerde gözlenen satın alma oranı {yuzde(istatistikler['visitor_type']['New_Visitor']['mean'])}, "
    f"geri dönenlerde {yuzde(istatistikler['visitor_type']['Returning_Visitor']['mean'])}. Bu fark ilişki gösterir; "
    "ziyaretçi türünün tek başına satın almaya neden olduğunu göstermez."
)
gorsel("product_pages.png", "Şekil 3. Ziyaret edilen ürün sayfalarının sonuca göre dağılımı.", 5.4)
paragraf(
    "Satın alma olan oturumlarda ziyaret edilen ürün sayfası medyanı 29, olmayanlarda 16'dır. "
    "Dağılımlar örtüştüğü için bu değişken tek başına kesin bir karar kuralı sunmaz. "
    "Veri %60 eğitim, %20 doğrulama ve %20 test olarak, sınıf oranları korunarak bölündü "
    f"({istatistikler['split']['train']:,}/{istatistikler['split']['validation']:,}/{istatistikler['split']['test']:,} kayıt). "
    "Sayısal değişkenler lojistik regresyon için yalnızca eğitim verisine uyarlanan StandardScaler ile "
    "ölçeklendirildi. Kategorik değişkenlere OneHotEncoder uygulandı; Random Forest'ta sayısal "
    "değişkenler olduğu gibi bırakıldı."
)

rapor.add_page_break()
baslik("3 Modeller ve değerlendirme")
paragraf(
    "Temel karşılaştırma için daima çoğunluk sınıfını söyleyen model kullanıldı. Lojistik regresyonda "
    "C=1, max_iter=2000 ve class_weight='balanced'; Random Forest'ta 200 ağaç, min_samples_leaf=5 ve "
    "class_weight='balanced_subsample' kullanıldı. Rastgelelik tohumu 42 olarak sabitlendi. "
    "Her model yalnızca eğitim kümesinde eğitildi; model seçimi doğrulama kümesindeki F1 ile yapıldı. "
    "Test kümesi son değerlendirme için ayrıldı ve karar eşiği 0,5 tutuldu."
)
baslik("3.1 Test sonuçları", 2)
metrik_tablosu()
paragraf(
    f"Çoğunluk modeli {yuzde(temel_model['accuracy'])} accuracy ile hiçbir satın alma oturumunu yakalayamadı "
    f"(recall {yuzde(temel_model['recall'])}). Random Forest'ın doğrulama F1 puanı "
    f"{yuzde(modeller['Random Forest']['validation']['f1'])} olup lojistik regresyonun "
    f"{yuzde(modeller['Lojistik regresyon']['validation']['f1'])} değerinden yüksektir. "
    f"Testte Random Forest precision {yuzde(orman_modeli['precision'])}, recall {yuzde(orman_modeli['recall'])} "
    f"ve F1 {yuzde(orman_modeli['f1'])} üretti."
)
gorsel("confusion_matrix.png", "Şekil 4. Seçilen modelin test kümesindeki karmaşıklık matrisi.", 4.7)
paragraf(
    f"Testte {orman_modeli['confusion_matrix'][1][0]} satın alma oturumu kaçırıldı (yanlış negatif); "
    f"{orman_modeli['confusion_matrix'][0][1]} satın alma olmayan oturum ise yanlış alarm verdi "
    "(yanlış pozitif). Hangi hata türünün daha pahalı olduğu kullanım senaryosuna göre belirlenmelidir."
)

rapor.add_page_break()
baslik("4 Hata analizi ve sınırlılıklar")
paragraf(
    "Seçilen model 382 gerçek satın alma oturumunun 293'ünü yakaladı, 89'unu kaçırdı. "
    "Yanlış negatifler olası hedeflemelerin kaçırılmasına, 201 yanlış pozitif ise gereksiz "
    "müdahalelere yol açabilir. Kaçırılan satın almalarda PageValues medyanı 0 iken doğru yakalananlarda "
    f"{hata_profilleri['true_positive']['median_page_values']:.1f}; yanlış pozitiflerde ürün sayfası medyanı "
    f"{hata_profilleri['false_positive']['median_product_pages']:.0f} iken doğru negatiflerde "
    f"{hata_profilleri['true_negative']['median_product_pages']:.0f}'tür. Bunlar hata örüntüleridir, "
    "nedensel açıklama değildir."
)
paragraf(
    f"PageValues çıkarılarak yapılan duyarlılık kontrolünde Random Forest test F1 değeri "
    f"{yuzde(duyarlilik_sonucu['f1'])}, recall değeri {yuzde(duyarlilik_sonucu['recall'])} oldu. "
    "Bu belirgin düşüş, modelin bu değişkene güçlü biçimde dayandığını gösterir. PageValues'ın "
    "tahmin anında nasıl hesaplandığı ve gelecek işlem bilgisini içerip içermediği doğrulanmadan "
    "gerçek zamanlı başarı iddiası kurulamaz."
)
paragraf(
    "Veri tek bir dönem ve kaynaktan gelir; rastgele bölme farklı dönemlerdeki performansı ölçmez. "
    "Oturum kimliği ve zaman damgası bulunmadığı için kullanıcı bazlı veya ileriye dönük zaman "
    "ayrımı yapılamadı. Bu çalışma, hazır veri üzerinde bir yöntem karşılaştırmasıdır; canlı ortam "
    "etkisini doğrulamaz."
)
baslik("4.1 Değişkenlerin kullanımı", 2)
tablo = rapor.add_table(rows=1, cols=3)
tablo.autofit = False
ozellik_genislikleri = [1.2, 2.7, 2.9]
for cell, value, width in zip(tablo.rows[0].cells, ["Değişken", "Verideki anlamı", "Uygulama açısından dikkat"], ozellik_genislikleri):
    cell.text = value
    cell.width = Inches(width)
    zemin_rengi(cell, "DCE8F2")
    kenarliklar(cell)
    for run in cell.paragraphs[0].runs:
        run.bold = True
        run.font.size = Pt(8.6)
ozellik_satirlari = [
    ("PageValues", "Ziyaret edilen sayfaların ortalama işlem değeri", "Tahmin anındaki hesaplama yöntemi doğrulanmalı"),
    ("ProductRelated", "Ziyaret edilen ürün sayfası sayısı", "Oturumun hangi anına kadar sayıldığı sabitlenmeli"),
    ("VisitorType", "Yeni veya geri dönen ziyaretçi", "Gözlenen oran farkı nedensellik olarak okunmamalı"),
    ("Month", "Ziyaret ayı", "Yeni dönemlerde dağılım değişimi izlenmeli"),
]
for i, row in enumerate(ozellik_satirlari):
    cells = tablo.add_row().cells
    for cell, value, width in zip(cells, row, ozellik_genislikleri):
        cell.text = value
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        kenarliklar(cell)
        if i % 2:
            zemin_rengi(cell, "F6F9FC")
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            for run in para.runs:
                run.font.size = Pt(8.5)
baslik("5 Sonuç ve geliştirme önerileri")
paragraf(
    f"Random Forest, ayrılmış test kümesinde {yuzde(orman_modeli['f1'])} F1 ile iki öğrenen model içinde daha "
    "iyi sonuç verdi. Öncelikli sonraki adım, tahmin anında gerçekten mevcut olan özellikleri "
    "tanımlamak ve modeli bu kısıtla yeniden eğitmektir. Ardından zamana göre ayrı bir test "
    "kümesi ve yanlış pozitif/negatif maliyetlerine göre karar eşiği çalışması yapılmalıdır."
)
baslik("Kaynak ve yeniden üretim", 2)
paragraf(
    "Sakar, C. ve Kastro, Y. (2018). Online Shoppers Purchasing Intention Dataset. "
    "UCI Machine Learning Repository. https://doi.org/10.24432/C5F88Q "
    "(CC BY 4.0)."
)
paragraf(
    "Analiz kodu ve raporun yeniden üretim adımları: " + KOD_DEPOSU
)

rapor_yolu = KOK / "E_ticaret_satin_alma_tahmini_raporu.docx"
rapor.save(rapor_yolu)
print(rapor_yolu)
